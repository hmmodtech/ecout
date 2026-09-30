import os
import io
import time
import wave
import queue
import threading
from datetime import datetime
from tempfile import NamedTemporaryFile
import customtkinter as ctk
import pyaudiowpatch as pyaudio
import numpy as np
from faster_whisper import WhisperModel

try:
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

RECORDINGS_DIR = "recordings"
os.makedirs(RECORDINGS_DIR, exist_ok=True)

# تحميل موديل Whisper المخصص للإنجليزية (دقيق وسريع)
print("[Meeting Engine] Loading Faster-Whisper model...")
whisper_model = WhisperModel("base.en", device="cpu", compute_type="int8")

class MeetingAudioEngine:
    def __init__(self):
        self.audio_q = queue.Queue()
        self.p = pyaudio.PyAudio()
        self.running = False

    def get_loopback_device(self):
        wasapi_info = self.p.get_host_api_info_by_type(pyaudio.paWASAPI)
        default_speakers = self.p.get_device_info_by_index(wasapi_info["defaultOutputDevice"])
        if not default_speakers["isLoopbackDevice"]:
            for loopback in self.p.get_loopback_device_info_generator():
                if default_speakers["name"] in loopback["name"]:
                    return loopback
            return self.p.get_default_wasapi_loopback()
        return default_speakers

    def record_stream(self, device_index, tag):
        stream = self.p.open(
            format=pyaudio.paInt16,
            channels=1,
            rate=16000,
            input=True,
            input_device_index=device_index,
            frames_per_buffer=1024
        )
        while self.running:
            try:
                data = stream.read(1024, exception_on_overflow=False)
                self.audio_q.put((tag, data))
            except Exception:
                pass
        stream.stop_stream()
        stream.close()

    def start(self):
        self.running = True
        # مايك المستخدم
        mic_idx = self.p.get_default_input_device_info()["index"]
        threading.Thread(target=self.record_stream, args=(mic_idx, "You"), daemon=True).start()

        # سماعات الجهاز / Teams / Meet
        spk_idx = self.get_loopback_device()["index"]
        threading.Thread(target=self.record_stream, args=(spk_idx, "Teams"), daemon=True).start()

    def stop(self):
        self.running = False

class MeetingApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Abujami Meeting Logger (Paragraph Mode)")
        self.geometry("860x560")
        self.configure(fg_color="#0d1117")

        self.engine = MeetingAudioEngine()
        self.is_recording = False
        self.meeting_paragraphs = []
        self.current_speaker = None
        self.current_text_buffer = ""
        self.last_audio_time = time.time()

        # الشريط العلوي
        self.top_bar = ctk.CTkFrame(self, fg_color="#161b22", height=45)
        self.top_bar.pack(fill="x", padx=10, pady=8)

        self.btn_toggle = ctk.CTkButton(
            self.top_bar, text="⏺ Start Recording Meeting", 
            fg_color="#238636", hover_color="#2ea043", 
            font=("Segoe UI", 12, "bold"), command=self.toggle_recording
        )
        self.btn_toggle.pack(side="left", padx=10, pady=6)

        self.lbl_status = ctk.CTkLabel(self.top_bar, text="Ready to record meeting as paragraphs", font=("Segoe UI", 11), text_color="#8b949e")
        self.lbl_status.pack(side="left", padx=10)

        # صندوق النص المجمع بالفقرات
        self.textbox = ctk.CTkTextbox(self, font=("Segoe UI", 13), fg_color="#161b22", text_color="#f0f6fc", wrap="word")
        self.textbox.pack(fill="both", expand=True, padx=10, pady=(0, 10))

    def toggle_recording(self):
        if not self.is_recording:
            self.is_recording = True
            self.meeting_paragraphs.clear()
            self.current_speaker = None
            self.current_text_buffer = ""
            self.textbox.delete("0.0", "end")
            self.btn_toggle.configure(text="⏹ Stop & Save Word", fg_color="#da3633", hover_color="#b62324")
            self.lbl_status.configure(text="● Recording meeting into continuous paragraphs...", text_color="#f85149")
            self.engine.start()
            threading.Thread(target=self.process_meeting_audio, daemon=True).start()
        else:
            self.is_recording = False
            self.engine.stop()
            self.btn_toggle.configure(text="⏺ Start Recording Meeting", fg_color="#238636", hover_color="#2ea043")
            self.lbl_status.configure(text="Saving Word document...", text_color="#e8c84a")
            self.flush_current_paragraph()
            self.save_to_word()

    def process_meeting_audio(self):
        active_tag = None
        buffer = bytearray()
        silence_threshold = 2.0  # انتظار ثانيتين من السكوت لتأكيد انتهاء العبارة

        while self.is_recording:
            try:
                tag, chunk = self.engine.audio_q.get(timeout=0.2)
                raw = np.frombuffer(chunk, dtype=np.int16)
                energy = np.abs(raw).mean()

                if energy > 250:
                    if active_tag != tag and len(buffer) > 16000:
                        self.transcribe_and_append(active_tag, bytes(buffer))
                        buffer = bytearray()
                    active_tag = tag
                    buffer.extend(chunk)
                    self.last_audio_time = time.time()
                else:
                    if len(buffer) > 16000 and (time.time() - self.last_audio_time > silence_threshold):
                        self.transcribe_and_append(active_tag, bytes(buffer))
                        buffer = bytearray()
                        active_tag = None
            except queue.Empty:
                if len(buffer) > 16000 and (time.time() - self.last_audio_time > silence_threshold):
                    self.transcribe_and_append(active_tag, bytes(buffer))
                    buffer = bytearray()
                    active_tag = None

    def transcribe_and_append(self, tag, audio_bytes):
        with NamedTemporaryFile(suffix=".wav", delete=False) as f:
            wav_path = f.name
        with wave.open(wav_path, "wb") as wf:
            wf.setnchannels(1)
            wf.setsampwidth(2)
            wf.setframerate(16000)
            wf.writeframes(audio_bytes)

        try:
            segments, _ = whisper_model.transcribe(wav_path, language="en", beam_size=2)
            text = " ".join([s.text.strip() for s in segments if s.text.strip()])
            if text and len(text) > 4:
                self.after(0, self.append_paragraph_text, tag, text)
        finally:
            if os.path.exists(wav_path):
                os.remove(wav_path)

    def append_paragraph_text(self, speaker, text):
        now_str = datetime.now().strftime("%H:%M:%S")
        # إذا كان المتحدث هو نفسه -> دمج النص كفقرة مستمرة دون تكرار الوقت
        if self.current_speaker == speaker:
            self.current_text_buffer += f" {text}"
            self.textbox.insert("end", f" {text}")
        else:
            # انتقال لمتحدث جديد -> سطر جديد + وسم الوقت والمتحدث مرة واحدة فقط
            self.flush_current_paragraph()
            self.current_speaker = speaker
            self.current_text_buffer = text
            spk_label = "[Remote / Teams]" if speaker == "Teams" else "[Candidate / You]"
            header = f"\n\n[{now_str}] {spk_label}:\n"
            self.textbox.insert("end", header + text)
            self.meeting_paragraphs.append({"speaker": spk_label, "time": now_str, "text": text})

        self.textbox.see("end")

    def flush_current_paragraph(self):
        if self.meeting_paragraphs and self.current_text_buffer:
            self.meeting_paragraphs[-1]["text"] = self.current_text_buffer

    def save_to_word(self):
        if not self.meeting_paragraphs:
            self.lbl_status.configure(text="No speech recorded.", text_color="#8b949e")
            return

        ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        filepath = os.path.join(RECORDINGS_DIR, f"Meeting_Transcript_{ts}.docx")

        if DOCX_AVAILABLE:
            doc = docx.Document()
            title = doc.add_heading("Conference & Meeting Proceedings", level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph(f"Generated on: {datetime.now().strftime('%A, %B %d, %Y - %I:%M %p')}")
            doc.add_paragraph("─" * 45)

            for p in self.meeting_paragraphs:
                p_head = doc.add_paragraph()
                r_time = p_head.add_run(f"[{p['time']}] ")
                r_time.font.color.rgb = RGBColor(120, 120, 120)
                r_spk = p_head.add_run(f"{p['speaker']}:")
                r_spk.bold = True
                r_spk.font.color.rgb = RGBColor(31, 110, 235) if "Teams" in p["speaker"] else RGBColor(35, 134, 54)

                # كتابة الفقرة كاملة متصلة
                p_body = doc.add_paragraph(p["text"].strip())
                p_body.paragraph_format.space_after = Pt(10)

            doc.save(filepath)
            self.lbl_status.configure(text="Word Saved & Opened!", text_color="#7ee787")
            try:
                os.startfile(os.path.abspath(filepath))
            except Exception:
                pass
        else:
            txt_path = os.path.join(RECORDINGS_DIR, f"Meeting_Transcript_{ts}.txt")
            with open(txt_path, "w", encoding="utf-8") as f:
                for p in self.meeting_paragraphs:
                    f.write(f"[{p['time']}] {p['speaker']}:\n{p['text']}\n\n")
            self.lbl_status.configure(text="Saved as TXT (python-docx not found)", text_color="#7ee787")

if __name__ == "__main__":
    app = MeetingApp()
    app.mainloop()