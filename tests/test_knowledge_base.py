import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from knowledge_base import (KnowledgeBase, chunk_passages, extract_text_from_file,  # noqa: E402
                            parse_qa_pairs)

LABELED = """
Interview Question Bank

Q1: Tell me about yourself.
A1: I am a security professional with 8 years of experience in hostile environments,
including Iraq and South Sudan. I specialise in risk assessment and team leadership.

Q2: How do you handle a hostile checkpoint?
A2: Stay calm, keep hands visible, follow the instructions of the person in charge,
never argue, and inform the operations room as soon as it is safe.

Question 3: What is your biggest weakness?
Answer: I used to take on too much myself; I now delegate and use checklists.
"""

NUMBERED = """
1. Why do you want to work with our organisation?
Because your mandate matches my experience in humanitarian security, and I believe
in your approach to community acceptance.

2. Describe a time you dealt with an aggressive crowd.
In 2021 a distribution turned tense. I withdrew the team to the vehicles, kept
communication with the crowd leaders, and de-escalated through the local elders.

3. What would you do if your convoy was ambushed?
Drive through if possible, otherwise stop, take cover away from the vehicles and
report to the ops room.
"""

DIALOGUE = """
Interviewer: Can you walk me through your medical training?
Candidate: I hold a current first aid certificate with tourniquet and haemorrhage control.
Interviewer: And what would you do first at a mine incident?
Candidate: Stop, do not move, warn others, and retrace my exact steps only if safe.
"""

ARABIC = """
س: حدثني عن نفسك
ج: أنا مسؤول أمن بخبرة ثماني سنوات.
س: كيف تتعامل مع نقطة تفتيش عدائية؟
ج: أبقى هادئاً وأتبع التعليمات.
"""

MARKDOWN = """
# Q&A

## What are your salary expectations?
I am flexible and focused on the role; the market range for this position is fine with me.

## How do you prioritise competing tasks?
Safety first, then time critical deliverables, then everything else. I communicate early.
"""


def test_parse_labeled_pairs():
    pairs = parse_qa_pairs(LABELED, "qa.txt")
    qs = [p.question for p in pairs]
    assert len(pairs) == 3
    assert qs[0] == "Tell me about yourself."
    assert "hostile checkpoint" in qs[1]
    assert pairs[1].answer.startswith("Stay calm")
    assert "delegate" in pairs[2].answer
    assert all(p.category == "QA" for p in pairs)


def test_parse_numbered_pairs():
    pairs = parse_qa_pairs(NUMBERED, "qa.txt")
    assert len(pairs) == 3
    assert pairs[0].question.startswith("Why do you want")
    assert "community acceptance" in pairs[0].answer
    assert pairs[2].question.startswith("What would you do if your convoy")
    assert "take cover" in pairs[2].answer


def test_parse_dialogue_pairs():
    pairs = parse_qa_pairs(DIALOGUE, "qa.txt")
    assert len(pairs) == 2
    assert "medical training" in pairs[0].question
    assert "tourniquet" in pairs[0].answer
    assert "mine incident" in pairs[1].question


def test_parse_arabic_pairs():
    pairs = parse_qa_pairs(ARABIC, "qa.txt")
    assert len(pairs) == 2
    assert pairs[1].question.startswith("كيف")


def test_parse_markdown_headings():
    pairs = parse_qa_pairs(MARKDOWN, "qa.md")
    assert len(pairs) == 2
    assert pairs[0].question == "What are your salary expectations?"
    assert "market range" in pairs[0].answer


def test_parse_returns_empty_for_prose():
    prose = "HEAT training covers many topics. Safety is important. Teams should plan ahead."
    assert parse_qa_pairs(prose) == []


def test_chunk_passages_keeps_headings():
    text = """# Checkpoints

Approach slowly and stop where indicated. Keep your hands visible at all times. Have your documents ready.
Do not take photographs. If the situation escalates, comply and report afterwards.

# Kidnapping

The first 24 hours are the most dangerous. Stay calm and do not resist physically. Try to humanise yourself to the captors.
""" * 3
    passages = chunk_passages(text, "heat.txt", "HEAT", max_chars=200)
    assert len(passages) >= 4
    titles = {p.title for p in passages}
    assert "Checkpoints" in titles and "Kidnapping" in titles
    assert all(len(p.text) <= 260 for p in passages)
    assert all(p.category == "HEAT" for p in passages)


@pytest.mark.skipif(not __import__("knowledge_base").DOCX_AVAILABLE, reason="python-docx missing")
def test_docx_table_and_paragraph_extraction(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Interview Q&A", level=1)
    d.add_paragraph("Q: What motivates you?")
    d.add_paragraph("A: Making a difference in difficult places.")
    table = d.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Question"
    table.cell(0, 1).text = "Answer"
    table.cell(1, 0).text = "How do you manage stress?"
    table.cell(1, 1).text = "Routine, exercise and talking to my team."
    table.cell(2, 0).text = "What is an IED?"
    table.cell(2, 1).text = "An improvised explosive device."
    path = tmp_path / "qa_bank.docx"
    d.save(str(path))

    text = extract_text_from_file(str(path))
    assert "## Interview Q&A" in text
    assert "Q: How do you manage stress?" in text
    pairs = parse_qa_pairs(text, "qa_bank.docx")
    questions = [p.question for p in pairs]
    assert "What motivates you?" in questions
    assert "How do you manage stress?" in questions
    assert "What is an IED?" in questions


def test_knowledge_base_categorises_and_indexes(tmp_path):
    ctx = tmp_path / "context"
    ctx.mkdir()
    (ctx / "qa_bank.txt").write_text(LABELED, encoding="utf-8")
    (ctx / "heat_security.txt").write_text(
        "# Checkpoints\n\nApproach slowly and stop where indicated. Keep hands visible. Have documents ready. "
        "Never argue with armed personnel and report the incident afterwards to the operations room.\n\n"
        "Q: What is the golden rule at a checkpoint?\nA: Comply, stay calm and report.\n"
        "Q: Who decides on route selection?\nA: The security focal point with the driver.\n"
        "Q: When do you use the radio check-in?\nA: Every 30 minutes on the move.\n",
        encoding="utf-8")
    (ctx / "cv_candidate.txt").write_text("John Doe – Security Officer. 8 years experience.", encoding="utf-8")
    (ctx / "jd_target.txt").write_text("Role: Field Security Advisor. Requirements: HEAT certified.", encoding="utf-8")
    (ctx / "notes.txt").write_text("Random extra notes about the organisation.", encoding="utf-8")

    kb = KnowledgeBase(str(ctx))
    stats = kb.load()
    assert len(stats.files["QA"]) == 1
    assert len(stats.files["HEAT"]) == 1
    assert len(stats.files["CV"]) == 1
    assert len(stats.files["JD"]) == 1
    assert len(stats.files["EXTRA"]) == 1
    assert stats.qa_pairs["QA"] == 3
    assert stats.qa_pairs["HEAT"] == 3
    assert len(kb.pairs_by_category("QA")) == 3
    assert len(kb.pairs_by_category("HEAT")) == 3
    assert any(p.category == "HEAT" for p in kb.passages)
    assert "John Doe" in kb.raw["CV"]
    assert "Q&A" not in stats.summary() or "QA" in stats.summary()
    hint = kb.vocabulary_hint()
    assert isinstance(hint, str)


def test_category_for_is_case_insensitive_and_keyword_based():
    cases = {
        "Security Interview Questions.docx": "QA",
        "Q&A Bank.docx": "QA",
        "qa.txt": "QA",
        "HEAT notes.pdf": "HEAT",
        "Hostile environment training.pdf": "HEAT",
        "Abujami_CV.pdf": "CV",
        "JD - Team Leader.docx": "JD",
        "notes.txt": "EXTRA",
        "extra_interview tips.txt": "EXTRA",   # forced prefix wins
    }
    for name, cat in cases.items():
        assert KnowledgeBase.category_for(name) == cat, name
